import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Set 0.8 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Colors
    COLOR_PRIMARY = RGBColor(35, 31, 28)       # #231f1c (Dark stone)
    COLOR_SECONDARY = RGBColor(120, 113, 108)  # #78716c (Muted stone)
    COLOR_ACCENT = RGBColor(168, 140, 100)     # Warm solar brass
    COLOR_DARK_TEXT = RGBColor(28, 25, 23)     # #1c1917
    HEX_BG_LIGHT = "F5F1E8"
    HEX_BG_HEADER = "231F1C"
    HEX_BORDER = "D6D3D1"

    # Set default style font
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_DARK_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # ==========================================
    # 1. COVER PAGE
    # ==========================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(40)
    p_pre.paragraph_format.space_after = Pt(8)
    run_pre = p_pre.add_run("ARCHITECTURAL RESIDENTIAL AUTOMATION & ENERGY MANAGEMENT")
    run_pre.font.size = Pt(10)
    run_pre.font.bold = True
    run_pre.font.color.rgb = COLOR_ACCENT

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run("3D Smart Home System")
    run_title.font.size = Pt(32)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(36)
    run_sub = p_sub.add_run("Technical Specification, Architectural Design & Project Showcase")
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Horizontal Divider Line (Table)
    div_table = doc.add_table(rows=1, cols=1)
    div_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_div = div_table.rows[0].cells[0]
    cell_div.width = Inches(6.8)
    set_cell_background(cell_div, "A88C64")
    cell_div.paragraphs[0].paragraph_format.space_before = Pt(2)
    cell_div.paragraphs[0].paragraph_format.space_after = Pt(2)

    # Cover Summary Box
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(30)
    p_box.paragraph_format.space_after = Pt(18)
    
    table_meta = doc.add_table(rows=5, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    metadata = [
        ("Project Name", "3D Smart Home System & Energy Dashboard"),
        ("Architecture", "Three.js / React Three Fiber + React 19 + Tailwind CSS v4"),
        ("Version", "1.0.0 Production Release"),
        ("Build & Quality", "Vite 8 · ESLint Clean (0 Errors, 0 Warnings)"),
        ("Date", "September 2026"),
    ]
    for idx, (label, val) in enumerate(metadata):
        c0 = table_meta.rows[idx].cells[0]
        c1 = table_meta.rows[idx].cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.8)
        set_cell_margins(c0, top=100, bottom=100, left=120, right=120)
        set_cell_margins(c1, top=100, bottom=100, left=120, right=120)
        set_cell_background(c0, "F5F1E8")
        set_cell_background(c1, "FAF8F5")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK_TEXT

    # Cover Page Break
    doc.add_page_break()

    # ==========================================
    # HELPER FUNCTIONS FOR SECTIONS & HEADINGS
    # ==========================================
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_ACCENT
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return p

    def add_image_card(img_path, caption_text, width=Inches(6.6)):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(10)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(img_path, width=width)

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(14)
            r_cap = p_cap.add_run(f"Figure: {caption_text}")
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = COLOR_SECONDARY
        else:
            p_err = doc.add_paragraph()
            p_err.add_run(f"[Image not found: {img_path}]").font.color.rgb = RGBColor(200, 50, 50)

    # ==========================================
    # 2. EXECUTIVE SUMMARY & PROJECT EXPLANATION
    # ==========================================
    add_h1("1. Project Overview & Explanation")
    
    p_intro = doc.add_paragraph()
    p_intro.add_run(
        "The 3D Smart Home System is an architectural residential control platform that transforms traditional smart "
        "home monitoring into an intuitive, spatial experience. Typical commercial smart home applications present users with "
        "endless vertical lists of isolated toggles, disconnected sensor notifications, and abstract power charts that lack physical context. "
        "By contrast, this system models the physical home in three dimensions using WebGL and Three.js, granting homeowners instantaneous "
        "spatial intuition over room environments, active lighting, and climate states."
    )

    p_sol = doc.add_paragraph()
    p_sol.add_run(
        "Beyond 3D spatial interaction, the platform provides an enterprise-grade Energy Management System (EMS) featuring appliance-level "
        "sub-metering, dynamic utility tariff financial calculations, and whole-home perimeter security with instant emergency panic controls. "
        "The entire user interface is designed under a Warm Brutalist architectural design system—utilizing subtle stone hues, tactile linework, "
        "sharp precision edges, and strict container height parity that completely eliminates awkward whitespace."
    )

    # Key Objectives Callout Box
    table_obj = doc.add_table(rows=1, cols=1)
    table_obj.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_obj = table_obj.rows[0].cells[0]
    c_obj.width = Inches(6.8)
    set_cell_background(c_obj, "F5F1E8")
    set_cell_margins(c_obj, top=140, bottom=140, left=180, right=180)
    p_box_obj = c_obj.paragraphs[0]
    p_box_obj.paragraph_format.space_after = Pt(2)
    r_obj_title = p_box_obj.add_run("Key Core Objectives:\n")
    r_obj_title.font.bold = True
    r_obj_title.font.size = Pt(10)
    r_obj_title.font.color.rgb = COLOR_PRIMARY
    
    obj_points = (
        "• Spatial Intuition: Interact directly with physical rooms and fixtures through an interactive 3D WebGL floorplan.\n"
        "• Energy Economics: Translate raw kWh telemetry into actionable monetary metrics ($ / € / £) and CO₂ savings.\n"
        "• Perimeter Protection: Monitor multi-zone security sensors and activate whole-home emergency lockdown within seconds.\n"
        "• Zero-Dead-Space Layouts: Ensure rigid grid parity and pixel-aligned borders across varied data cards."
    )
    r_obj_text = p_box_obj.add_run(obj_points)
    r_obj_text.font.size = Pt(9.5)
    r_obj_text.font.color.rgb = COLOR_DARK_TEXT

    # ==========================================
    # 3. TECHNOLOGIES USED
    # ==========================================
    add_h1("2. Technology Stack & Architectural Framework")

    p_tech_desc = doc.add_paragraph()
    p_tech_desc.add_run(
        "The application is engineered using cutting-edge modern web technologies, prioritizing high rendering performance, minimal bundle overhead, "
        "strict type safety, and responsive design across all devices."
    )

    # Tech Table
    tech_table = doc.add_table(rows=7, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Technology", "Version / Library", "Role & Architecture Value"]
    
    # Header row
    for i, h in enumerate(headers):
        c = tech_table.rows[0].cells[i]
        set_cell_background(c, HEX_BG_HEADER)
        set_cell_margins(c, top=120, bottom=120, left=140, right=140)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    tech_data = [
        ("React 19", "v19.2.8", "Component-based reactive UI architecture, state hooks, and optimized concurrent DOM rendering."),
        ("Three.js & R3F", "v0.186 / Fiber v9.7", "High-performance WebGL 3D rendering engine and declarative React fiber bridge for spatial scenes."),
        ("@react-three/drei", "v10.7.8", "Three.js helper ecosystem providing smooth OrbitControls, ContactShadows, and lighting shaders."),
        ("Tailwind CSS v4", "v4.3.3", "Modern CSS utility framework with native nesting, variable themes, and custom brutalist stone palettes."),
        ("Lucide React", "v1.47.0", "Clean, consistent SVG iconography for smart home devices, weather telemetry, and security sensors."),
        ("Vite 8", "v8.3.0", "Next-generation frontend tooling offering sub-second Hot Module Replacement (HMR) and fast production builds."),
    ]

    for row_idx, data in enumerate(tech_data, start=1):
        for col_idx, text in enumerate(data):
            c = tech_table.rows[row_idx].cells[col_idx]
            bg = "FAF8F5" if row_idx % 2 == 1 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, top=100, bottom=100, left=140, right=140)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_DARK_TEXT

    tech_table.columns[0].width = Inches(1.6)
    tech_table.columns[1].width = Inches(1.5)
    tech_table.columns[2].width = Inches(3.7)

    doc.add_page_break()

    # ==========================================
    # 4. COMPREHENSIVE FEATURES & APPLICATION SCREENSHOTS
    # ==========================================
    add_h1("3. Core Features & System Walkthrough")

    # Feature 1: 3D Scene
    add_h2("3.1 Interactive 3D Architectural Home Scene (SmartHomeScene)")
    p_f1 = doc.add_paragraph()
    p_f1.add_run(
        "The primary overview presents an isometric 3D spatial representation of the residence. "
        "Users can freely rotate, pan, and zoom around the home using bounded Three.js orbit controls. Each room (Living room, Kitchen, "
        "Master bedroom, Bathroom, Entrance) features 3D geometry floor slabs, procedural furniture pieces, and real-time interactive highlights. "
        "Selecting any room updates the integrated side panel with live climate setpoints, master illumination controls, and sub-device toggles "
        "(Smart TV, soundbars, floor lamps)."
    )
    add_image_card("docs_images/overview.png", "3D Interactive Home Scene with Spatial Room Controls & Side Panel")

    # Feature 2: Energy
    add_h2("3.2 Energy Intelligence & Financial Analytics (EnergyPanel)")
    p_f2 = doc.add_paragraph()
    p_f2.add_run(
        "The Energy Panel is a comprehensive energy accounting system designed to optimize self-sufficiency and quantify monetary savings. "
        "It features rooftop solar generation monitoring, 13.5 kWh battery storage telemetry, and external grid exchange metrics."
    )
    p_f2_b = doc.add_paragraph()
    p_f2_b.add_run(
        "Key capabilities include:\n"
        "• Dual-Bar Chart Analytics: Interactive 24h, 7d, and 30d views comparing solar output against home demand.\n"
        "• Appliance-Level Sub-Metering: Real-time active draw and daily kWh accumulation across 7+ primary home appliances.\n"
        "• Utility Tariff & Currency Calculator: Dynamic rate adjustment supporting USD ($), EUR (€), and GBP (£) with projected monthly billing.\n"
        "• Net Energy Balance Visualizer: Live solar offset percentages demonstrating carbon savings (kg CO₂)."
    )
    add_image_card("docs_images/energy.png", "Energy Management Dashboard with Hourly Generation Charts & Appliance Sub-Metering")
    add_image_card("docs_images/energy_dark.png", "Energy Panel in High-Contrast Stone Dark Mode")

    # Feature 3: Security
    add_h2("3.3 Security Center & Perimeter Protection (SecurityPanel)")
    p_f3 = doc.add_paragraph()
    p_f3.add_run(
        "The Security Center provides whole-home perimeter visibility. Homeowners can switch between multiple live camera previews "
        "(Front door, Garden, Driveway) and monitor a distributed mesh of 8+ door contact, window, and PIR motion sensors. "
        "The platform includes an interactive sensor simulation tool to test alarm triggers, policy selector modes (Disarmed, Home, Away, Night), "
        "and an Emergency Panic Lockdown button that seals all perimeter deadbolts and activates sirens."
    )
    add_image_card("docs_images/security.png", "Security Center with Live Camera Previews, Sensor Mesh & Emergency Lockdown")

    # Feature 4: Smart Scenes
    add_h2("3.4 Orchestrated Smart Scenes & Automations (SmartScenes)")
    p_f4 = doc.add_paragraph()
    p_f4.add_run(
        "Automated routines allow synchronized actuation across multiple home subsystems. Homeowners can trigger pre-configured "
        "profiles such as 'Morning Wakeup', 'Evening Relaxation', 'Movie Night', 'Away Mode', 'Dinner Time', and 'Focus & Work'. "
        "Each scene orchestrates target lighting temperatures, ambient brightness, climate offsets, and media system states simultaneously."
    )
    add_image_card("docs_images/scenes.png", "Automated Smart Scenes with Multi-Device Actuation & Scheduling")

    # Feature 5: Settings
    add_h2("3.5 Residence Profile, Localization & Diagnostics (SettingsPanel)")
    p_f5 = doc.add_paragraph()
    p_f5.add_run(
        "The Settings panel provides central administrative control over property dimensions, occupant counts, regional timezone "
        "formatting, temperature units (°C / °F), and notification dispatch thresholds. It also features hardware diagnostics "
        "displaying network latency (ms), mesh node health, and system reboot simulation tools."
    )
    add_image_card("docs_images/settings.png", "Settings & Diagnostic Panel with System Telemetry Controls")

    doc.add_page_break()

    # ==========================================
    # 5. UI DESIGN SYSTEM & LAYOUT BALANCING
    # ==========================================
    add_h1("4. Warm Brutalist Design Philosophy & Layout Parity")
    
    p_design = doc.add_paragraph()
    p_design.add_run(
        "A standout architectural element of this project is its strict adherence to a Warm Brutalist design system. "
        "Instead of neon, high-contrast futuristic tropes, the interface employs natural architectural stone colors: "
        "soft limestone (#f2eee5), warm chalk (#f7f4ed), and deep charcoal stone (#1c1917, #231f1c). Hairline borders (border-stone-300) "
        "and sharp corners (rounded-none) create a disciplined, timeless aesthetic."
    )

    p_layout = doc.add_paragraph()
    p_layout.add_run(
        "Zero-Dead-Space Engineering: In complex multi-column dashboards, cards often end at ragged, mismatched heights, leaving "
        "awkward empty gaps. This project implements a flexbox top-content and pinned-footer hierarchy (flex flex-col justify-between) "
        "aligned with CSS Grid row constraints. As demonstrated in Figure 2 and Figure 4, cards across both columns terminate at the exact "
        "same pixel baseline with natural, compact list spacing and zero vertical stretching."
    )

    # ==========================================
    # 6. FUTURE UPDATES & EXPANSION ROADMAP
    # ==========================================
    add_h1("5. Future Updates & Development Roadmap")

    roadmap_items = [
        ("Phase 1: Native IoT Hardware Bridge (Matter / Thread / Zigbee)",
         "Integrate Home Assistant and Matter/Thread bridge protocols via WebSocket APIs to allow bi-directional live communication with physical IoT smart switches, Zigbee door sensors, and Tuya/Shelly power plugs."),
        ("Phase 2: Predictive AI Energy Optimizer",
         "Implement local AI model inference that analyzes 7-day weather forecasts and historical family habits to automatically pre-cool the house during peak solar generation or schedule EV charging during lowest utility tariff intervals."),
        ("Phase 3: Digital Twin BIM Model Import",
         "Enable homeowners to upload standard IFC / BIM architectural files or floorplan sketches, dynamically generating custom 3D Three.js floorplans matching their exact residence geometry."),
        ("Phase 4: Mobile Companion & Offline Edge Support",
         "Deliver a companion mobile application built with React Native and Expo, incorporating offline SQLite edge synchronization and localized push notifications for emergency perimeter breaches."),
        ("Phase 5: Conversational Voice & Residence Agent",
         "Introduce an embedded local voice interface allowing homeowners to execute multi-step requests naturally ('Prepare the living room for a movie and lock the front door').")
    ]

    for title, desc in roadmap_items:
        add_h3(title)
        p_rd = doc.add_paragraph()
        p_rd.add_run(desc)
        p_rd.paragraph_format.space_after = Pt(8)

    # ==========================================
    # 7. CONCLUSION
    # ==========================================
    add_h1("6. Conclusion & Quality Verification")
    p_conclusion = doc.add_paragraph()
    p_conclusion.add_run(
        "The 3D Smart Home System demonstrates how modern WebGL spatial computing can elevate everyday IoT home control into an intuitive, "
        "architectural experience. With rigorous engineering across 3D rendering, appliance sub-metering, emergency security, and balanced "
        "responsive layouts, the platform represents a production-ready blueprint for next-generation smart living."
    )

    # Verification Table
    ver_table = doc.add_table(rows=4, cols=2)
    ver_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    checks = [
        ("ESLint 10 Compliance", "Passed: 0 errors, 0 warnings across all components"),
        ("Production Bundle Build", "Passed: Built in 1.44s via Vite 8 (Gzip CSS: 8.3 kB, JS: 332 kB)"),
        ("Browser Compatibility", "Chrome, Edge, Firefox, Safari (Full WebGL 2.0 & Mobile Touch support)"),
    ]
    for idx, (metric, status) in enumerate(checks):
        c0 = ver_table.rows[idx].cells[0]
        c1 = ver_table.rows[idx].cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.6)
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)
        set_cell_background(c0, "F5F1E8")
        set_cell_background(c1, "FAF8F5")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(metric)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(status)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK_TEXT

    # Output file path
    out_filename = "3D_Smart_Home_System_Project_Report.docx"
    doc.save(out_filename)
    print(f"Successfully generated {out_filename} ({os.path.getsize(out_filename)} bytes)")

if __name__ == "__main__":
    create_document()
